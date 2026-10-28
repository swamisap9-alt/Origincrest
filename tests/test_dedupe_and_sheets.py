from datetime import date

import openpyxl

from origincrest_agent.dedupe import DedupeIndex, company_domain, core_name, norm_phone, primary_key
from origincrest_agent.workbook import WorkbookStore, build_dedupe_index


def test_company_name_normalisation():
    assert core_name("Seaway Logistics Pty Ltd") == core_name("SEAWAY LOGISTICS")
    assert core_name("The Bean & Co.") == "bean"


def test_domains_and_phones():
    assert company_domain("https://www.example.com.au/contact") == "example.com.au"
    assert company_domain("facebook.com/somecafe") == ""
    assert company_domain("https://shop.myshopify.com") == "shop.myshopify.com"
    assert norm_phone("+61 8 8100 0000") == norm_phone("(08) 8100 0000")


def test_dedupe_index_matches_variants():
    index = DedupeIndex()
    index.add("row 1", company="Seven Seeds Coffee Roasters Pty Ltd", website="https://sevenseeds.com.au")
    assert index.match(company="Seven Seeds") is not None  # loose name
    assert index.match(company="Totally Different", website="http://www.sevenseeds.com.au/shop") is not None
    assert index.match(company="Other", email="wholesale@sevenseeds.com.au") is not None  # same domain
    assert index.match(company="Market Lane Coffee") is None


def test_free_mail_is_not_a_company_domain():
    assert primary_key("Tiny Roaster", "", "tinyroaster@gmail.com") == "email:tinyroaster@gmail.com"
    index = DedupeIndex()
    index.add("row", company="Tiny Roaster", email="tinyroaster@gmail.com")
    assert index.match(company="Another Roaster", email="another@gmail.com") is None


def test_header_detection_and_column_mapping(cfg):
    store = WorkbookStore("rfq", cfg.workbook_path("rfq"), cfg).load()
    view = store.view("RFQ Tracker")
    assert view.header_row == 3
    assert view.columns["company"] == 1
    assert view.columns["contact_name"] == 2
    assert view.columns["email"] == 3
    assert view.columns["last_contacted"] == 7
    assert "agent_notes" not in view.columns  # user's "Notes" is never taken over
    assert [r.text("company") for r in view.rows()] == ["Seaway Logistics", "Fastfreight Customs", "Harbour 3PL"]


def test_add_columns_extends_excel_table(cfg):
    store = WorkbookStore("rfq", cfg.workbook_path("rfq"), cfg).load()
    view = store.view("RFQ Tracker")
    row = next(view.rows())
    row.set("next_action_date", date(2026, 11, 1))
    view.append({"company": "New Co", "status": "Open"})
    store.save(cfg.data_dir / "backups")

    wb = openpyxl.load_workbook(cfg.workbook_path("rfq"))
    ws = wb["RFQ Tracker"]
    assert ws.cell(3, 9).value == "Next action date"
    assert ws.cell(4, 9).value.date() == date(2026, 11, 1)
    assert ws.cell(7, 1).value == "New Co"
    table = ws.tables["RFQs"]
    assert table.ref == "A3:I7"
    assert [c.name for c in table.tableColumns][-1] == "Next action date"
    assert list((cfg.data_dir / "backups").rglob("*.xlsx"))


def test_dedupe_index_spans_all_tabs(cfg):
    stores = {n: WorkbookStore(n, cfg.workbook_path(n), cfg).load() for n in ("rfq", "master")}
    index = build_dedupe_index(stores)
    assert index.match(company="Bean There Roasters Pty Ltd") is not None
    assert index.match(company="X", website="paf.com.au") is not None
    assert index.match(company="Y", email="quotes@seaway.com.au") is not None


def test_new_tab_has_full_header(cfg):
    store = WorkbookStore("master", cfg.workbook_path("master"), cfg).load()
    view = store.get_or_create("New Leads - Adelaide", "buyer")
    assert view.kind == "buyer"
    for key in ("company", "email", "source_url", "status", "fit_score", "next_action_date"):
        assert view.has(key)


def test_name_matches_need_more_than_one_short_word():
    index = DedupeIndex()
    index.add("row 1", company="Hills Cafe Group")
    assert index.match(company="Hills Roasting Co") is None
    index.add("row 2", company="Allpress Espresso Pty Ltd")
    assert index.match(company="Allpress Coffee Roastery") is not None  # one distinctive long word


def test_same_name_different_websites_are_different_companies():
    index = DedupeIndex()
    index.add("row 1", company="Bean Bar", website="beanbar.com.au")
    assert index.match(company="Bean Bar", website="beanbarperth.com.au") is None
    assert index.match(company="Bean Bar") is not None
    assert index.match(company="Bean Bar Cafe", website="https://www.beanbar.com.au") is not None


def test_keep_awake_asks_windows_not_to_sleep(monkeypatch):
    from origincrest_agent import utils

    calls = []

    class FakeKernel32:
        def SetThreadExecutionState(self, flags):
            calls.append(flags)

    monkeypatch.setattr(utils, "_kernel32", lambda: FakeKernel32())
    with utils.keep_awake():
        assert calls == [0x80000001]  # ES_CONTINUOUS | ES_SYSTEM_REQUIRED while running
    assert calls[-1] == 0x80000000  # back to normal afterwards

    monkeypatch.setattr(utils, "_kernel32", lambda: None)  # not Windows: no-op
    with utils.keep_awake():
        pass
