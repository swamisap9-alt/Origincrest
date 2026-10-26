"""Test fixtures: sample workbooks shaped like real ones, a fake Claude agent and a fake mailbox."""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from origincrest_agent.config import DEFAULTS, Config, _deep_merge
from origincrest_agent.llm import ResearchResult
from origincrest_agent.mailer import MailInfo
from origincrest_agent.runtime import Report, RunContext
from origincrest_agent.state import State
from origincrest_agent.verify import Verifier

TODAY = date(2026, 10, 26)


def make_rfq_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RFQ Tracker"
    ws["A1"] = "OriginCrest freight RFQ tracker"
    headers = ["Company Name", "Contact Person", "Email Address", "Phone", "Service", "Status",
               "Date Contacted", "Notes"]
    for c, h in enumerate(headers, start=1):
        ws.cell(3, c, h)
    rows = [
        ["Seaway Logistics", "Jo Bloggs", "quotes@seaway.com.au", "08 8100 0000", "FCL Belawan-Adelaide", "Open",
         TODAY - timedelta(days=7), "Asked for 1x20ft"],
        ["Fastfreight Customs", "", "info@fastfreight.com.au", "", "Customs clearance", "Open",
         TODAY - timedelta(days=2), ""],
        ["Harbour 3PL", "", "hello@harbour3pl.com.au", "", "Pallet storage", "Quoted",
         TODAY - timedelta(days=20), "Quote received"],
    ]
    for r, values in enumerate(rows, start=4):
        for c, v in enumerate(values, start=1):
            ws.cell(r, c, v)
    ws.add_table(Table(displayName="RFQs", ref="A3:H6"))
    wb.save(path)


def make_master_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All contacts"
    for c, h in enumerate(["Company", "Business Type", "Website", "Email", "State"], start=1):
        ws.cell(1, c, h)
    ws.append(["Port Adelaide Freight Pty Ltd", "3PL", "https://www.paf.com.au", "", "SA"])
    ws.append(["Bean There Roasters", "Roasters", "beanthere.com.au", "info@beanthere.com.au", "VIC"])
    ws.append(["Pacific Green Imports", "Importers", "pacificgreen.com.au", "", "NSW"])
    ws.append(["Bulk Bean Traders", "Wholesalers & Traders", "", "", "QLD"])
    prospects = wb.create_sheet("New Prospects Oct-26")
    for c, h in enumerate(["Company", "Category", "State"], start=1):
        prospects.cell(1, c, h)
    prospects.append(["Hills Cafe Group", "Cafe groups", "SA"])
    prospects.append(["Bean There Roasters", "Roasters", "VIC"])
    wb.save(path)


@pytest.fixture
def cfg(tmp_path: Path) -> Config:
    rfq, master = tmp_path / "rfq.xlsx", tmp_path / "master.xlsx"
    make_rfq_workbook(rfq)
    make_master_workbook(master)
    data = _deep_merge(DEFAULTS, {
        "data_dir": str(tmp_path / "data"),
        "workbooks": {"rfq": {"path": str(rfq)}, "master": {"path": str(master)}},
        "excel": {"lock_wait_minutes": 0},
        "sender": {"name": "Swami K", "business_name": "OriginCrest", "address": "1 Test St, Adelaide SA 5000",
                   "email": "hello@origincrest.com.au", "phone": "0400 000 000"},
        "offer": {"summary": "OriginCrest imports traceable Indonesian single-origin green coffee."},
    })
    return Config(data)


class FakeAgent:
    def __init__(self) -> None:
        self.research_queue: list[ResearchResult] = []
        self.structured_queue: list[dict] = []
        self.calls: list[tuple[str, str]] = []
        self.usage = {"requests": 0}

    def research(self, *, system, prompt, record_tool, effort, web_searches, web_fetches) -> ResearchResult:
        self.calls.append(("research", prompt))
        return self.research_queue.pop(0)

    def structured(self, *, system, prompt, schema, effort) -> dict:
        self.calls.append(("structured", prompt))
        if self.structured_queue:
            return self.structured_queue.pop(0)
        return {"subject": "Indonesian green for your range", "body": "Hi there,\n\nShort note.\n\nKind regards,"}


def mail(folder: str, when: datetime, from_addr: str, to: str, subject: str, text: str = "",
         message_id: str = "", auto: bool = False) -> MailInfo:
    return MailInfo(folder, "1", when, from_addr, [to], subject, message_id or f"<{abs(hash((subject, when)))}@x>",
                    "", auto, text)


class FakeMailbox:
    def __init__(self) -> None:
        self.drafts: list = []
        self.sent: list[MailInfo] = []
        self.inbox: list[MailInfo] = []
        self.connected = False

    def connect(self) -> None:
        self.connected = True

    def close(self) -> None:
        self.connected = False

    def append_draft(self, msg) -> None:
        self.drafts.append(msg)

    def draft_waiting(self, message_id: str = "", to_addr: str = "") -> bool:
        return any((message_id and m["Message-ID"] == message_id) or (to_addr and to_addr in m["To"])
                   for m in self.drafts)

    def messages_to(self, address: str, since: date) -> list[MailInfo]:
        hits = [m for m in self.sent if address in m.to_addrs and m.when.date() >= since]
        return sorted(hits, key=lambda m: m.when, reverse=True)

    def messages_from(self, address: str, since: date, domain: str = "") -> list[MailInfo]:
        hits = [m for m in self.inbox if (m.from_addr == address or (domain and m.from_addr.endswith("@" + domain)))
                and m.when.date() >= since and not m.is_bounce]
        return sorted(hits, key=lambda m: m.when, reverse=True)

    def bounces_for(self, address: str, since: date) -> list[MailInfo]:
        return [m for m in self.inbox if m.is_bounce and address in m.text and m.when.date() >= since]


class FakeVerifier(Verifier):
    """Verifier with canned pages instead of network fetches."""

    def __init__(self, pages: dict[str, str] | None = None):
        super().__init__(enabled=True)
        self.pages = pages or {}

    def fetch(self, url: str) -> str | None:
        return self.pages.get(url)


@pytest.fixture
def agent() -> FakeAgent:
    return FakeAgent()


@pytest.fixture
def mailbox() -> FakeMailbox:
    return FakeMailbox()


@pytest.fixture
def make_ctx(cfg, agent, mailbox):
    def _make(workflow: str = "test", pages: dict[str, str] | None = None, dry_run: bool = False) -> RunContext:
        return RunContext(cfg=cfg, state=State(cfg.data_dir / "state.db"), today=TODAY,
                          report=Report(workflow, TODAY), dry_run=dry_run, agent_factory=lambda: agent,
                          mailbox_factory=lambda: mailbox, verifier=FakeVerifier(pages))
    return _make


def utc(d: date, hour: int = 9) -> datetime:
    return datetime(d.year, d.month, d.day, hour, tzinfo=timezone.utc)
