from datetime import timedelta

import openpyxl

from origincrest_agent.followup import run_followup
from origincrest_agent.workbook import WorkbookStore

from .conftest import TODAY, mail, utc


def rfq_rows(cfg):
    ws = openpyxl.load_workbook(cfg.workbook_path("rfq"))["RFQ Tracker"]
    headers = [c.value for c in ws[3]]
    return {r[0].value: dict(zip(headers, (c.value for c in r))) for r in ws.iter_rows(min_row=4) if r[0].value}


def sent_rfq(mailbox, days_ago=7):
    msg = mail("Sent", utc(TODAY - timedelta(days=days_ago)), "hello@origincrest.com.au", "quotes@seaway.com.au",
               "RFQ: 1x20ft green coffee Belawan to Adelaide", text="Hi Jo, could you quote 1x20ft...",
               message_id="<rfq-1@origincrest.com.au>")
    mailbox.sent.append(msg)
    return msg


def test_quote_followup_threaded(cfg, agent, mailbox, make_ctx):
    sent_rfq(mailbox)
    ctx = make_ctx("followup")
    assert run_followup(ctx) == 1
    msg = mailbox.drafts[0]
    assert msg["To"] == "Jo Bloggs <quotes@seaway.com.au>"
    assert msg["Subject"] == "Re: RFQ: 1x20ft green coffee Belawan to Adelaide"
    assert msg["In-Reply-To"] == "<rfq-1@origincrest.com.au>"
    assert "requests like this" in msg.get_content()  # short RFQ opt-out line, not the marketing footer
    assert "> Hi Jo, could you quote" in msg.get_content()

    rows = rfq_rows(cfg)
    seaway = rows["Seaway Logistics"]
    assert seaway["Date Contacted"].date() == TODAY  # existing column reused
    assert seaway["Next action date"].date() == TODAY + timedelta(days=5)
    assert seaway["Follow-ups drafted"] == 1
    assert "Follow-up #1 drafted" in seaway["Agent notes"]
    assert seaway["Notes"] == "Asked for 1x20ft"  # user's notes untouched
    assert rows["Fastfreight Customs"]["Agent notes"] is None  # only 2 days, not due
    assert rows["Harbour 3PL"]["Agent notes"] is None  # not Open

    # Same day again: nothing new.
    assert run_followup(make_ctx("followup")) == 0
    assert len(mailbox.drafts) == 1


def test_reply_found_stops_followup(cfg, agent, mailbox, make_ctx):
    sent = sent_rfq(mailbox)
    mailbox.inbox.append(mail("INBOX", sent.when + timedelta(days=2), "jo@seaway.com.au", "hello@origincrest.com.au",
                              "Re: RFQ: 1x20ft green coffee", text="Quote attached."))
    run_followup(make_ctx("followup"))
    assert mailbox.drafts == []
    seaway = rfq_rows(cfg)["Seaway Logistics"]
    assert seaway["Status"] == "Replied - review"
    assert seaway["Last reply"].date() == (sent.when + timedelta(days=2)).date()


def test_auto_reply_is_not_a_reply(cfg, agent, mailbox, make_ctx):
    sent = sent_rfq(mailbox)
    mailbox.inbox.append(mail("INBOX", sent.when + timedelta(hours=1), "quotes@seaway.com.au",
                              "hello@origincrest.com.au", "Automatic reply: RFQ", auto=True))
    assert run_followup(make_ctx("followup")) == 1
    assert "Auto-reply only" in rfq_rows(cfg)["Seaway Logistics"]["Agent notes"]


def test_unsubscribe_is_honoured(cfg, agent, mailbox, make_ctx):
    sent = sent_rfq(mailbox)
    mailbox.inbox.append(mail("INBOX", sent.when + timedelta(days=1), "quotes@seaway.com.au",
                              "hello@origincrest.com.au", "Re: RFQ", text="Please remove me from your list."))
    ctx = make_ctx("followup")
    run_followup(ctx)
    assert mailbox.drafts == []
    assert rfq_rows(cfg)["Seaway Logistics"]["Status"] == "Unsubscribed - do not contact"
    assert ctx.state.is_suppressed("quotes@seaway.com.au")


def test_lead_intro_sent_becomes_open_and_uses_10_days(cfg, agent, mailbox, make_ctx):
    # A lead in the master workbook's new-lead tab is followed up too (follow-ups otherwise read only the
    # RFQ workbook).
    store = WorkbookStore("master", cfg.workbook_path("master"), cfg).load()
    view = store.get_or_create("New Leads - Adelaide", "buyer")
    view.append({"company": "Hills Roasting Co", "email": "info@hillsroasting.com.au",
                 "status": "Intro drafted - to review", "draft_created": TODAY - timedelta(days=12),
                 "record_type": "Lead"})
    store.save(cfg.data_dir / "backups")
    mailbox.sent.append(mail("Sent", utc(TODAY - timedelta(days=11)), "hello@origincrest.com.au",
                             "info@hillsroasting.com.au", "Indonesian green for your range", text="Hi there"))
    run_followup(make_ctx("followup"))
    ws = openpyxl.load_workbook(cfg.workbook_path("master"))["New Leads - Adelaide"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, (c.value for c in ws[2])))
    assert row["Status"] == "Open"
    # 11 days since the intro went out > 10-day lead window, so a follow-up is drafted with the unsubscribe footer.
    lead_draft = next(m for m in mailbox.drafts if "hillsroasting" in m["To"])
    assert "unsubscribe" in lead_draft.get_content()
    assert lead_draft["List-Unsubscribe"]
    assert row["Next action date"].date() == TODAY + timedelta(days=10)


def test_max_followups_flagged(cfg, agent, mailbox, make_ctx):
    sent_rfq(mailbox)
    store = WorkbookStore("rfq", cfg.workbook_path("rfq"), cfg).load()
    next(store.view("RFQ Tracker").rows()).set("followup_count", 3)
    store.save(cfg.data_dir / "backups")
    ctx = make_ctx("followup")
    run_followup(ctx)
    assert mailbox.drafts == []
    assert "Flagged - max follow-ups reached" in ctx.report.sections


def test_failed_sheet_write_does_not_duplicate_draft(cfg, agent, mailbox, make_ctx):
    sent_rfq(mailbox)
    rfq = cfg.workbook_path("rfq")
    owner = rfq.parent / ("~$" + rfq.name)
    owner.write_text("Excel")
    run_followup(make_ctx("followup"))
    assert len(mailbox.drafts) == 1
    assert rfq_rows(cfg)["Seaway Logistics"]["Status"] == "Open"  # not written: Excel had it open
    owner.unlink()
    run_followup(make_ctx("followup"))
    assert len(mailbox.drafts) == 1  # no second draft
    seaway = rfq_rows(cfg)["Seaway Logistics"]
    assert seaway["Follow-ups drafted"] == 1
    assert seaway["Next action date"].date() == TODAY + timedelta(days=5)


def test_sent_rfq_becomes_open_with_5_day_chase(cfg, agent, mailbox, make_ctx):
    store = WorkbookStore("rfq", cfg.workbook_path("rfq"), cfg).load()
    store.get_or_create("New Freight-3PL - Other States", "freight").append({
        "company": "Sydney Bonded Warehousing", "email": "rfq@sbw.com.au", "status": "RFQ drafted - to review",
        "draft_created": TODAY - timedelta(days=3), "record_type": "Quote"})
    store.save(cfg.data_dir / "backups")
    mailbox.sent.append(mail("Sent", utc(TODAY - timedelta(days=2)), "hello@origincrest.com.au", "rfq@sbw.com.au",
                             "RFQ: pallet storage Sydney", text="Please quote"))
    ctx = make_ctx("followup")
    run_followup(ctx)
    ws = openpyxl.load_workbook(cfg.workbook_path("rfq"))["New Freight-3PL - Other States"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, (c.value for c in ws[2])))
    assert row["Status"] == "Open"
    assert row["Next action date"].date() == TODAY + timedelta(days=3)  # sent 2 days ago + 5
    assert not [m for m in mailbox.drafts if "sbw" in m["To"]]  # not due yet


def add_contacts_tab(cfg):
    """A tab like the real "contacts" one: no Status column, your own Notes / Date sent / Followup date."""
    wb = openpyxl.load_workbook(cfg.workbook_path("rfq"))
    ws = wb.create_sheet("contacts")
    ws.append(["Company", "Contact", "Email", "Phone", "Services", "Notes", "Date sent", "Followup date"])
    for row in [
        ["Smith Channon", "Ann Lee", "ops@smithchannon.com.au", "08 8200 1111", "Customs brokerage", None,
         TODAY - timedelta(days=8), None],
        ["Pakval", None, "info@pakval.com.au", None, "Fumigation", None, TODAY - timedelta(days=8),
         TODAY + timedelta(days=2)],
        ["CTI Interstate", "Ben", "sales@cti.com.au", "08 8300 2222", "Interstate transport", "Need to be called",
         TODAY - timedelta(days=6), None],
        ["Adelaide Devanning", None, None, "0411 111 111", "Devanning", "needs to be called re pallet rates",
         TODAY - timedelta(days=9), None],
        ["Quick Quote Co", None, "rfq@quickquote.com.au", None, None, "Quote received 20/10",
         TODAY - timedelta(days=20), None],
        ["Not Yet Sent Pty", None, "hello@notyet.com.au", None, None, None, None, None],
        ["Early Chase Freight", None, "rfq@earlychase.com.au", None, None, None, TODAY - timedelta(days=2),
         TODAY - timedelta(days=1)],
    ]:
        ws.append(row)
    wb.save(cfg.workbook_path("rfq"))


def contacts_rows(cfg):
    ws = openpyxl.load_workbook(cfg.workbook_path("rfq"))["contacts"]
    headers = [c.value for c in ws[1]]
    return {r[0].value: dict(zip(headers, (c.value for c in r))) for r in ws.iter_rows(min_row=2) if r[0].value}


def call_aware_agent(agent):
    def structured(*, system, prompt, schema, effort):
        agent.calls.append(("structured", prompt))
        points = ["Check the customs quote is coming", "Ask about lead time"] if "call_points" in prompt else []
        return {"subject": "Following up on our RFQ", "body": "Hi,\n\nJust following up on our request.\n\nThanks,",
                "call_points": points}
    agent.structured = structured


def test_contacts_tab_columns_are_read_not_taken_over(cfg):
    add_contacts_tab(cfg)
    view = WorkbookStore("rfq", cfg.workbook_path("rfq"), cfg).load().view("contacts")
    assert view.columns["date_sent"] == 7
    assert view.columns["next_action_date"] == 8
    assert view.columns["user_notes"] == 6
    assert not view.has("status") and not view.has("last_contacted")


def test_contacts_tab_followups_use_date_sent_followup_date_and_call_notes(cfg, agent, mailbox, make_ctx):
    add_contacts_tab(cfg)
    call_aware_agent(agent)
    ctx = make_ctx("followup")
    run_followup(ctx)
    drafts = {m["To"]: m for m in mailbox.drafts}

    # Date sent 8 days ago, no Followup date: due after 5 days.
    smith = drafts["Ann Lee <ops@smithchannon.com.au>"]
    assert smith["Subject"] == "Following up on our RFQ"
    assert "CALL FIRST" not in smith.get_content()
    # Your Followup date (yesterday) wins over date sent + 5.
    assert "rfq@earlychase.com.au" in drafts
    # Not due (Followup date ahead), closed by the notes, or never sent.
    assert not any(a in to for to in drafts for a in ("pakval", "quickquote", "notyet"))

    # "Need to be called" with an email: flagged follow-up with the phone number and call points.
    cti = drafts["Ben <sales@cti.com.au>"]
    assert cti["Subject"] == "[CALL FIRST] Following up on our RFQ"
    text = cti.get_content()
    assert text.index("CALL FIRST") < text.index("Just following up")
    assert "Your note: Need to be called" in text and "08 8300 2222" in text
    assert "- Check the customs quote is coming" in text
    assert 'delete this block and "[CALL FIRST]"' in text

    # No email: a call reminder addressed to you, with no marketing footer.
    reminder = drafts["Swami K <hello@origincrest.com.au>"]
    assert reminder["Subject"] == "[CALL] Adelaide Devanning: follow-up #1 due"
    assert "0411 111 111" in reminder.get_content() and "pallet rates" in reminder.get_content()
    assert "unsubscribe" not in reminder.get_content().lower() and not reminder["List-Unsubscribe"]
    assert "Call needed - follow-up drafted with [CALL FIRST]" in ctx.report.sections
    assert "Call reminders drafted (addressed to you)" in ctx.report.sections

    rows = contacts_rows(cfg)
    assert rows["Smith Channon"]["Date sent"].date() == TODAY - timedelta(days=8)  # yours, never overwritten
    assert rows["Smith Channon"]["Last contacted"].date() == TODAY
    assert rows["Smith Channon"]["Followup date"].date() == TODAY + timedelta(days=5)
    assert rows["Smith Channon"]["Notes"] is None
    assert rows["CTI Interstate"]["Last contacted"] is None  # you're calling, not emailing
    assert rows["CTI Interstate"]["Followup date"].date() == TODAY + timedelta(days=5)
    assert "Your notes say call" in rows["CTI Interstate"]["Agent notes"]
    assert rows["CTI Interstate"]["Notes"] == "Need to be called"
    assert rows["Pakval"]["Followup date"].date() == TODAY + timedelta(days=2)
    assert rows["Quick Quote Co"]["Agent notes"] is None

    # Next weekday run: nothing new until the next Followup date.
    count = len(mailbox.drafts)
    run_followup(make_ctx("followup"))
    assert len(mailbox.drafts) == count


def test_contacts_tab_reply_marks_row(cfg, agent, mailbox, make_ctx):
    add_contacts_tab(cfg)
    mailbox.inbox.append(mail("INBOX", utc(TODAY - timedelta(days=3)), "ann@smithchannon.com.au",
                              "hello@origincrest.com.au", "Re: RFQ customs", text="Quote attached."))
    run_followup(make_ctx("followup"))
    assert not [m for m in mailbox.drafts if "smithchannon" in m["To"]]
    row = contacts_rows(cfg)["Smith Channon"]
    assert row["Status"] == "Replied - review"  # added so later runs leave it alone
