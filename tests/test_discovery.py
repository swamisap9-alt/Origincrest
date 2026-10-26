import openpyxl

from origincrest_agent.discovery import profile_categories, run_discover
from origincrest_agent.llm import ResearchResult
from origincrest_agent.runtime import ChangeSet, commit
from origincrest_agent.workbook import WorkbookStore


def company(name, business_type="", **kw):
    base = {"company_name": name, "business_type": business_type, "website": "", "contact_name": "",
            "contact_role": "", "email": "", "phone": "", "city": "", "state": "", "adelaide_presence": False,
            "services": "", "coffee_evidence": "", "source_url": "", "notes": ""}
    base.update(kw)
    return base


PAGES = {
    "https://indofreight.com.au/contact": "<p>Sea freight desk: <a href='mailto:rfq@indofreight.com.au'>email</a>"
                                          " or call (08) 8447 1234</p>",
    "https://hillsroasting.com.au/contact": "<p>Wholesale: info@hillsroasting.com.au</p>",
    "https://quietroasters.com.au/contact": "<p>hello@quietroasters.com.au - no unsolicited emails please</p>",
}


def freight_adelaide():
    return ResearchResult(data={"companies": [
        company("Seaway Logistics Pty Ltd", "Freight forwarder", state="SA"),  # already in the RFQ tracker
        company("Indo Freight Lines", "Freight forwarder", website="indofreight.com.au", email="rfq@indofreight.com.au",
                phone="08 8447 1234", city="Port Adelaide", state="SA", adelaide_presence=True,
                coffee_evidence="Green coffee FCL from Belawan", source_url="https://indofreight.com.au/contact"),
    ]}, pages=[("https://indofreight.com.au/contact", PAGES["https://indofreight.com.au/contact"])])


def freight_other():
    return ResearchResult(data={"companies": [
        company("Sydney Bonded Warehousing", "3PL / warehouse", city="Botany", state="NSW"),
        company("National Freight Co", "Freight forwarder", city="Melbourne", state="VIC", adelaide_presence=True),
    ]})


def roasters_adelaide():
    return ResearchResult(data={"companies": [
        company("Hills Roasting Co", "Specialty roaster", website="hillsroasting.com.au",
                email="info@hillsroasting.com.au", city="Hahndorf", state="SA",
                source_url="https://hillsroasting.com.au/contact"),
        company("Quiet Roasters", "Specialty roaster", website="quietroasters.com.au",
                email="hello@quietroasters.com.au", city="Adelaide", state="SA",
                source_url="https://quietroasters.com.au/contact"),
    ]})


def roasters_other():
    return ResearchResult(data={"companies": [
        company("Bean There Roasters", "Specialty roaster", state="VIC"),  # already in All contacts
        company("Guessy Roasters", "Specialty roaster", website="guessy.com.au", email="john@guessy.com.au",
                city="Melbourne", state="VIC", source_url="https://guessy.com.au/about"),
    ]})


def rows_of(path, tab):
    ws = openpyxl.load_workbook(path)[tab]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, (c.value for c in r))) for r in ws.iter_rows(min_row=2)]


def test_master_categories_come_from_its_tabs(cfg):
    store = WorkbookStore("master", cfg.workbook_path("master"), cfg).load()
    notes = []
    cats = profile_categories(cfg, store, notes)
    assert [(c.label, c.kind) for c in cats] == [
        ("Cafe groups", "buyer"), ("Roasters", "buyer"),             # every category in New Prospects Oct-26
        ("Importers", "buyer"), ("Wholesalers & Traders", "buyer"),  # matching ones from All contacts (not 3PL)
    ]
    assert cats[1].examples == ("Bean There Roasters",)
    assert notes == []


def test_category_fallback_and_missing_tab(cfg):
    cfg.data["discovery"]["profiles"]["master"]["category_tabs"] = [
        {"tab": "new prospects oct 26"},  # matched despite case/punctuation
        {"tab": "All contacts", "match": ["exporter"], "fallback": "Importers, Roasters, Wholesalers & Traders"},
        {"tab": "Does not exist"},
    ]
    store = WorkbookStore("master", cfg.workbook_path("master"), cfg).load()
    notes = []
    labels = [c.label for c in profile_categories(cfg, store, notes)]
    assert labels == ["Cafe groups", "Roasters", "Importers", "Wholesalers", "Traders"]
    assert any("Does not exist" in n for n in notes)


def test_rfq_workbook_gets_freight_leads_split_by_region(cfg, agent, make_ctx):
    agent.research_queue = [freight_adelaide(), freight_other()]
    ctx = make_ctx("discover", pages=PAGES)
    assert run_discover(ctx, workbooks=["rfq"]) == 3

    rfq = cfg.workbook_path("rfq")
    adelaide = rows_of(rfq, "New Freight-3PL - Adelaide")
    assert [r["Company"] for r in adelaide] == ["Indo Freight Lines", "National Freight Co"]  # SA or Adelaide office
    indo = adelaide[0]
    assert indo["Email"] == "rfq@indofreight.com.au" and indo["Phone"] == "08 8447 1234"
    assert indo["Category"] == "Freight forwarder" and indo["Record type"] == "Quote"
    assert indo["Status"] == "New - to review"
    assert [r["Company"] for r in rows_of(rfq, "New Freight-3PL - Other States")] == ["Sydney Bonded Warehousing"]
    assert "Seaway" in " ".join(ctx.report.sections["Skipped duplicates"])
    # Nothing is written to the master workbook by the RFQ profile.
    assert "New Leads - Adelaide" not in openpyxl.load_workbook(cfg.workbook_path("master")).sheetnames
    prompt = agent.calls[0][1]
    assert "Freight forwarders, customs brokers & 3PLs" in prompt and "South Australia" in prompt


def test_master_workbook_gets_leads_for_its_categories(cfg, agent, make_ctx):
    agent.research_queue = [roasters_adelaide(), roasters_other()]
    ctx = make_ctx("discover", pages=PAGES)
    assert run_discover(ctx, workbooks=["master"], categories=["Roasters"]) == 3

    master = cfg.workbook_path("master")
    adelaide = rows_of(master, "New Leads - Adelaide")
    assert [r["Company"] for r in adelaide] == ["Hills Roasting Co", "Quiet Roasters"]
    assert adelaide[0]["Category"] == "Roasters" and adelaide[0]["Record type"] == "Lead"
    assert adelaide[0]["Email"] == "info@hillsroasting.com.au"
    assert adelaide[1]["Email"] is None and "no unsolicited" in adelaide[1]["Agent notes"]
    other = rows_of(master, "New Leads - Other States")
    assert [r["Company"] for r in other] == ["Guessy Roasters"]
    assert other[0]["Email"] is None and "Unverified email" in other[0]["Agent notes"]
    assert "Bean There Roasters" in agent.calls[0][1]  # sheet examples guide the search
    assert "Bean There" in " ".join(ctx.report.sections["Skipped duplicates"])

    # Running again finds only duplicates.
    agent.research_queue = [roasters_adelaide(), roasters_other()]
    assert run_discover(make_ctx("discover", pages=PAGES), workbooks=["master"], categories=["Roasters"]) == 0


def test_categories_rotate_and_caps_apply(cfg, agent, make_ctx):
    cfg.data["discovery"]["categories_per_run"] = 1
    cfg.data["discovery"]["max_new_rows_per_segment"] = 1
    agent.research_queue = [roasters_adelaide(), roasters_other()]
    ctx = make_ctx("discover", pages=PAGES)
    assert run_discover(ctx, workbooks=["master"]) == 2  # one category this week x two regions, 1 row each
    assert len(agent.calls) == 2
    assert any("cap of 1 reached" in line for line in ctx.report.sections["Caps"])


def test_rows_wait_while_workbook_is_open_in_excel(cfg, agent, make_ctx):
    rfq = cfg.workbook_path("rfq")
    owner = rfq.parent / ("~$" + rfq.name)
    owner.write_text("Excel")
    agent.research_queue = [freight_adelaide()]
    ctx = make_ctx("discover", pages=PAGES)
    run_discover(ctx, workbooks=["rfq"], regions=["adelaide"])
    assert "New Freight-3PL - Adelaide" not in openpyxl.load_workbook(rfq).sheetnames
    assert any("open in Excel" in line for line in ctx.report.sections["Not written"])

    owner.unlink()
    commit(make_ctx("flush"), ChangeSet())
    assert len(rows_of(rfq, "New Freight-3PL - Adelaide")) == 1
    commit(make_ctx("flush"), ChangeSet())  # nothing left queued: no duplicate rows
    assert len(rows_of(rfq, "New Freight-3PL - Adelaide")) == 1


def test_dry_run_writes_nothing(cfg, agent, make_ctx):
    agent.research_queue = [freight_adelaide()]
    ctx = make_ctx("discover", pages=PAGES, dry_run=True)
    run_discover(ctx, workbooks=["rfq"], regions=["adelaide"])
    assert "Would add" in ctx.report.sections
    assert "New Freight-3PL - Adelaide" not in openpyxl.load_workbook(cfg.workbook_path("rfq")).sheetnames
