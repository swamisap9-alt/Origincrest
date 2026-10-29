from datetime import timedelta

import openpyxl

from origincrest_agent.leads import run_leads
from origincrest_agent.llm import ResearchResult
from origincrest_agent.workbook import WorkbookStore

from .conftest import TODAY, mail, utc


LEADS_ADELAIDE = "New Leads - Adelaide"


def seed_buyers(cfg):
    rows = [
        {"company": "Hills Roasting Co", "category": "Roasters", "email": "info@hillsroasting.com.au",
         "website": "https://hillsroasting.com.au", "status": "New - to review",
         "source_url": "https://hillsroasting.com.au/contact", "added_by": "OriginCrest Agent", "record_type": "Lead"},
        {"company": "Low Fit Cafe", "category": "Cafe groups", "email": "hi@lowfit.com.au", "status": "New - to review",
         "added_by": "OriginCrest Agent", "record_type": "Lead"},
        {"company": "No Email Roasters", "category": "Roasters", "website": "https://noemail.com.au", "status": ""},
        {"company": "Existing Customer Roasters", "category": "Roasters", "email": "buy@customer.com.au",
         "status": "Customer"},
    ]
    store = WorkbookStore("master", cfg.workbook_path("master"), cfg).load()
    view = store.get_or_create(LEADS_ADELAIDE, "buyer")
    for r in rows:
        view.append(r)
    store.save(cfg.data_dir / "backups")


def research(name, score, email="", **kw):
    data = {"company_name": name, "website": "", "city": "Adelaide", "state": "SA",
            "what_they_roast": "Light-roast single origins, filter focus", "size_estimate": "Small",
            "size_evidence": "Two staff, 15kg roaster", "sourcing": "Buys from two importers",
            "current_indonesian_coffees": "", "wholesale": "Yes", "public_email": email,
            "public_email_source_url": "", "contact_name": "", "contact_role": "", "fit_score": score,
            "best_fit_origins": ["Sumatra", "Flores"], "fit_rationale": "Single-origin focus, no Indonesian lot.",
            "personalisation_hooks": [{"fact": "Sells a washed Ethiopian on filter", "source_url": "https://x"}],
            "red_flags": ""}
    data.update(kw)
    return ResearchResult(data=data)


def sheet(cfg, name="master", tab=LEADS_ADELAIDE):
    ws = openpyxl.load_workbook(cfg.workbook_path(name))[tab]
    headers = [c.value for c in ws[1]]
    return {r[0].value: dict(zip(headers, (c.value for c in r))) for r in ws.iter_rows(min_row=2)}


def test_research_score_and_intro(cfg, agent, mailbox, make_ctx):
    seed_buyers(cfg)
    # Order: rows you added yourself first, then the agent's finds.
    agent.research_queue = [research("No Email Roasters", 70), research("Hills Roasting Co", 82),
                            research("Low Fit Cafe", 30)]
    ctx = make_ctx("leads")
    assert run_leads(ctx) == 1
    # Only the new-lead tab is worked on: nothing in "All contacts" / "New Prospects Oct-26" is researched.
    assert len([c for c in agent.calls if c[0] == "research"]) == 3

    assert len(mailbox.drafts) == 1
    msg = mailbox.drafts[0]
    assert msg["To"] == "info@hillsroasting.com.au"
    assert msg["List-Unsubscribe"] == "<mailto:hello@origincrest.com.au?subject=unsubscribe>"
    body = msg.get_content()
    assert "Swami K" in body and "1 Test St" in body
    assert "unsubscribe" in body and "hillsroasting.com.au" in body

    rows = sheet(cfg)
    hills = rows["Hills Roasting Co"]
    assert hills["Status"] == "Intro drafted - to review"
    assert hills["Fit score (0-100)"] == 82
    assert hills["Best-fit origins"] == "Sumatra, Flores"
    assert hills["Coffee offering"] == "Light-roast single origins, filter focus"
    assert hills["Draft subject"] == "Indonesian green for your range"
    assert rows["Low Fit Cafe"]["Status"] == "Researched - low fit"
    assert rows["No Email Roasters"]["Status"] == "No public email"
    assert rows["Existing Customer Roasters"]["Status"] == "Customer"
    assert rows["Existing Customer Roasters"]["Fit score (0-100)"] is None

    # Nothing left to do on a second run: no new API calls, no new drafts.
    calls = len(agent.calls)
    assert run_leads(make_ctx("leads")) == 0
    assert len(agent.calls) == calls and len(mailbox.drafts) == 1


def test_rfq_drafted_for_new_freight_lead(cfg, agent, mailbox, make_ctx):
    store = WorkbookStore("rfq", cfg.workbook_path("rfq"), cfg).load()
    store.get_or_create("New Freight-3PL - Adelaide", "freight").append({
        "company": "Indo Freight Lines", "category": "Freight forwarder", "email": "rfq@indofreight.com.au",
        "city": "Port Adelaide", "state": "SA", "services": "Green coffee FCL from Belawan; customs clearance",
        "status": "New - to review", "record_type": "Quote", "added_by": "OriginCrest Agent"})
    store.save(cfg.data_dir / "backups")
    # Your earlier RFQ (to a contact in the RFQ tracker) is used as the reference for what to ask.
    mailbox.sent.append(mail("Sent", utc(TODAY - timedelta(days=7)), "hello@origincrest.com.au",
                             "quotes@seaway.com.au", "RFQ: 1x20ft green coffee Belawan to Adelaide",
                             text="Could you quote 1x20ft FCL, 19.2t green coffee in jute bags, FOB Belawan?"))
    agent.structured_queue = [{"subject": "RFQ: green coffee FCL Belawan to Port Adelaide",
                               "body": "Hi team,\n\nPlease quote [number of 20ft containers per year].\n\nKind regards,"}]
    ctx = make_ctx("leads")
    assert run_leads(ctx) == 1

    assert not [c for c in agent.calls if c[0] == "research"]  # no fit research for logistics providers
    prompt = agent.calls[0][1]
    assert "Green coffee FCL from Belawan" in prompt          # this company's own row
    assert "FCL Belawan-Adelaide" in prompt                   # what the RFQ tracker shows you ask for
    assert "19.2t green coffee in jute bags" in prompt        # your last RFQ email
    assert "quotes@seaway.com.au" not in prompt               # other providers' contact details stay out

    msg = mailbox.drafts[0]
    assert msg["To"] == "rfq@indofreight.com.au"
    assert msg["Subject"] == "RFQ: green coffee FCL Belawan to Port Adelaide"
    assert "requests like this" in msg.get_content()          # short opt-out line, not the marketing footer
    row = sheet(cfg, "rfq", "New Freight-3PL - Adelaide")["Indo Freight Lines"]
    assert row["Status"] == "RFQ drafted - to review"
    assert row["Draft subject"] == "RFQ: green coffee FCL Belawan to Port Adelaide"
    assert "Fill in the [bracketed] details" in row["Agent notes"]


def test_approved_status_overrides_low_score(cfg, agent, mailbox, make_ctx):
    seed_buyers(cfg)
    store = WorkbookStore("master", cfg.workbook_path("master"), cfg).load()
    view = store.view(LEADS_ADELAIDE)
    for row in view.rows():
        row.set("status", "Approved" if row.text("company") == "Low Fit Cafe" else "Customer")
    store.save(cfg.data_dir / "backups")
    agent.research_queue = [research("Low Fit Cafe", 30)]
    run_leads(make_ctx("leads"))
    assert [m["To"] for m in mailbox.drafts] == ["hi@lowfit.com.au"]


def test_no_drafts_until_sender_details_filled(cfg, agent, mailbox, make_ctx):
    seed_buyers(cfg)
    cfg.data["email"]["signature"] = "config"  # signature from config.yaml, but it isn't filled in
    cfg.data["sender"]["address"] = "<<EDIT ME: business postal address>>"
    cfg.data["sender"]["phone"] = ""
    agent.research_queue = [research("No Email Roasters", 70), research("Hills Roasting Co", 82),
                            research("Low Fit Cafe", 30)]
    ctx = make_ctx("leads")
    run_leads(ctx)
    assert mailbox.drafts == []
    assert "sender.address" in " ".join(ctx.report.sections["Errors"])
    assert sheet(cfg)["Hills Roasting Co"]["Researched on"] is None  # picked up again once fixed


def test_your_own_signature_is_used(cfg, agent, mailbox, make_ctx):
    """With no sender details in config.yaml, name, address and signature come from your Sent folder."""
    seed_buyers(cfg)
    cfg.data["sender"] = {"business_name": "OriginCrest"}
    signature = "Swami Kailasam\nDirector | OriginCrest\nM 0400 111 222 | www.origincrest.com.au"
    for days, (to, opener, signoff) in enumerate([("a@x.com.au", "Hi Jo,", "Kind regards,"),
                                                   ("b@y.com.au", "Hello Sam,", "Thanks,"),
                                                   ("c@z.com.au", "Hi team,", "Cheers,")], start=40):
        mailbox.sent.append(mail("Sent", utc(TODAY - timedelta(days=days)), "swami@origincrest.com.au", to,
                                 "Hello", text=f"{opener}\n\nSome message.\n\n{signoff}\n{signature}",
                                 from_name="Swami Kailasam"))
    agent.research_queue = [research("No Email Roasters", 70), research("Hills Roasting Co", 82),
                            research("Low Fit Cafe", 30)]
    ctx = make_ctx("leads")
    run_leads(ctx)
    msg = mailbox.drafts[0]
    assert msg["From"] == "Swami Kailasam <swami@origincrest.com.au>"
    body = msg.get_content()
    assert "Kind regards,\n\nSwami Kailasam\nDirector | OriginCrest\nM 0400 111 222" in body
    assert "unsubscribe" in body  # the opt-out footer is still added
    assert "your recent sent emails in 'Sent'" in " ".join(ctx.report.sections["Signature"])


def test_already_emailed_contact_goes_to_open(cfg, agent, mailbox, make_ctx):
    seed_buyers(cfg)
    sent_day = TODAY - timedelta(days=30)
    mailbox.sent.append(mail("Sent", utc(sent_day), "hello@origincrest.com.au", "info@hillsroasting.com.au",
                             "Coffee chat"))
    agent.research_queue = [research("No Email Roasters", 70), research("Hills Roasting Co", 82),
                            research("Low Fit Cafe", 30)]
    run_leads(make_ctx("leads"))
    hills = sheet(cfg)["Hills Roasting Co"]
    assert mailbox.drafts == []
    assert hills["Status"] == "Open"
    assert hills["Last contacted"].date() == sent_day


def test_draft_cap_defers_rest(cfg, agent, mailbox, make_ctx):
    seed_buyers(cfg)
    agent.research_queue = [research("No Email Roasters", 70, email=""), research("Hills Roasting Co", 82),
                            research("Low Fit Cafe", 90)]
    ctx = make_ctx("leads")
    run_leads(ctx, max_drafts=1)
    assert len(mailbox.drafts) == 1
    assert "Deferred (draft cap)" in ctx.report.sections
    agent.research_queue = []  # research is cached: the deferred row needs no new research
    run_leads(make_ctx("leads"), max_drafts=5)
    assert len(mailbox.drafts) == 2


def test_intro_is_written_from_the_contacts_own_row(cfg, agent, mailbox, make_ctx):
    store = WorkbookStore("master", cfg.workbook_path("master"), cfg).load()
    view = store.get_or_create(LEADS_ADELAIDE, "buyer")
    view.append({"company": "Hills Roasting Co", "email": "info@hillsroasting.com.au", "contact_name": "Mia Chen",
                 "status": "Reviewed", "talking_points": "Met Mia at the Adelaide coffee expo; she wants a Flores",
                 "agent_notes": "internal bookkeeping"})
    store.save(cfg.data_dir / "backups")
    agent.research_queue = [research("Hills Roasting Co", 80)]
    run_leads(make_ctx("leads"))
    prompt = next(p for kind, p in agent.calls if kind == "structured")
    assert "Met Mia at the Adelaide coffee expo" in prompt
    assert "Address it to Mia." in prompt
    assert "internal bookkeeping" not in prompt
    assert len(mailbox.drafts) == 1


def test_offer_section_is_optional(cfg, agent, mailbox, make_ctx):
    from origincrest_agent.compliance import config_sender_problems
    from origincrest_agent.prompts import DEFAULT_SUMMARY, offer_text

    cfg.data["offer"] = {"summary": "<<EDIT ME: e.g. OriginCrest imports...>>",
                         "origins": ["Sumatra - <<EDIT ME: e.g. Gayo>>", "Flores - <<EDIT ME>>"], "facts": [],
                         "call_to_action": ""}
    assert config_sender_problems(cfg) == []
    text = offer_text(cfg)
    assert DEFAULT_SUMMARY in text and "Origins offered: Sumatra; Flores" in text and "<<" not in text


def seed_freight(cfg, company="Indo Freight Lines", email="rfq@indofreight.com.au"):
    store = WorkbookStore("rfq", cfg.workbook_path("rfq"), cfg).load()
    store.get_or_create("New Freight-3PL - Adelaide", "freight").append({
        "company": company, "category": "Freight forwarder", "email": email, "city": "Port Adelaide",
        "state": "SA", "services": "Green coffee FCL", "status": "New - to review", "record_type": "Quote",
        "added_by": "OriginCrest Agent"})
    store.save(cfg.data_dir / "backups")


def test_redraft_after_you_deleted_the_draft(cfg, agent, mailbox, make_ctx):
    seed_freight(cfg)
    assert run_leads(make_ctx("leads")) == 1
    first_id = mailbox.drafts[0]["Message-ID"]
    assert run_leads(make_ctx("leads")) == 0  # normally never drafted twice

    # Still in Drafts: the redraft refuses, so you don't get two.
    ctx = make_ctx("redraft")
    assert run_leads(ctx, redraft=["indo freight", "Nobody Pty Ltd"]) == 0
    assert "Still in Drafts" in ctx.report.sections
    assert "Nobody Pty Ltd" in " ".join(ctx.report.sections["Not found"])

    mailbox.drafts.clear()  # you deleted it
    agent.structured_queue = [{"subject": "RFQ: FCL Belawan to Port Adelaide", "body": "Hi team,\n\nNew text.\n\nKind regards,"}]
    ctx = make_ctx("redraft")
    assert run_leads(ctx, redraft=["Indo Freight"]) == 1
    assert len(mailbox.drafts) == 1 and mailbox.drafts[0]["Message-ID"] != first_id
    row = sheet(cfg, "rfq", "New Freight-3PL - Adelaide")["Indo Freight Lines"]
    assert row["Status"] == "RFQ drafted - to review" and row["Draft subject"] == "RFQ: FCL Belawan to Port Adelaide"
    assert ctx.state.draft("domain:indofreight.com.au", "rfq", 1)["message_id"] == mailbox.drafts[0]["Message-ID"]


def test_redraft_skips_companies_you_already_emailed(cfg, agent, mailbox, make_ctx):
    seed_freight(cfg)
    run_leads(make_ctx("leads"))
    mailbox.drafts.clear()
    mailbox.sent.append(mail("Sent", utc(TODAY), "hello@origincrest.com.au", "rfq@indofreight.com.au", "RFQ"))
    ctx = make_ctx("redraft")
    assert run_leads(ctx, redraft=["Indo Freight Lines"]) == 0
    assert "Already contacted" in ctx.report.sections
    assert sheet(cfg, "rfq", "New Freight-3PL - Adelaide")["Indo Freight Lines"]["Status"] == "Open"
