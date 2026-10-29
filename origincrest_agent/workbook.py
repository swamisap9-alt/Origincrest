"""Excel access that adapts to each sheet's existing layout.

* The header row is detected (it doesn't have to be row 1) and existing columns
  are matched by name; missing columns are appended to the right.
* Every save is preceded by a timestamped backup, and nothing is written while
  Excel has the file open.
"""

from __future__ import annotations

import os
import shutil
import time
import warnings
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn

from .dedupe import DedupeIndex
from .schema import FIELD_BY_KEY, NEW_TAB_KEYS, NOT_FOR_EMAIL, category_kind, match_header, match_weak_header
from .utils import clean, fmt_date, log, to_date

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_SCAN_ROWS = 15


class WorkbookLocked(Exception):
    pass


def _value(ws, r: int, c: int) -> Any:
    """Read a cell without creating it (ws.cell() would add empty cells to the file)."""
    cell = ws._cells.get((r, c))
    return None if cell is None else cell.value


def excel_lock_reason(path: Path, stale_hours: float = 12) -> str | None:
    """Why the workbook can't be written right now, or None if it's free."""
    owners = {path.parent / ("~$" + path.name), path.parent / ("~$" + path.name[2:])}
    for owner in owners:
        if owner.exists() and (time.time() - owner.stat().st_mtime) / 3600 < stale_hours:
            return f"it is open in Excel ({owner.name} exists)"
    try:
        with open(path, "r+b"):
            pass
    except PermissionError:
        return "it is in use by another program"
    except FileNotFoundError:
        return None
    return None


def wait_until_unlocked(paths: list[Path], minutes: float) -> dict[Path, str]:
    """Wait up to `minutes` for the files to be closed; return the ones still locked."""
    deadline = time.monotonic() + minutes * 60
    while True:
        locked = {p: r for p in paths if (r := excel_lock_reason(p))}
        if not locked or time.monotonic() >= deadline:
            return locked
        log.info("Waiting for %s to be closed in Excel...", ", ".join(p.name for p in locked))
        time.sleep(30)


def _detect_header_row(ws) -> int | None:
    best_row, best_score = None, 0
    first_texty = None
    max_col = min(ws.max_column, 120)
    for r in range(1, min(ws.max_row, _SCAN_ROWS) + 1):
        score = texty = 0
        for c in range(1, max_col + 1):
            value = _value(ws, r, c)
            if isinstance(value, str) and value.strip():
                texty += 1
                if match_header(value) or match_weak_header(value):
                    score += 1
        if score > best_score:
            best_row, best_score = r, score
        if first_texty is None and texty >= 2:
            first_texty = r
    if best_row and best_score >= 2:
        return best_row
    return first_texty


class Row:
    def __init__(self, sheet: "SheetView", index: int):
        self.sheet = sheet
        self.index = index

    def get(self, key: str) -> Any:
        return self.sheet.get(self.index, key)

    def text(self, key: str) -> str:
        return clean(self.get(key))

    def set(self, key: str, value: Any) -> None:
        self.sheet.set(self.index, key, value)

    def context(self) -> dict[str, str]:
        return self.sheet.row_context(self.index)

    def last_contact_date(self) -> date | None:
        """The later of "Last contacted" and your own "Date sent"."""
        dates = [d for d in (to_date(self.get("last_contacted")), to_date(self.get("date_sent"))) if d]
        return max(dates) if dates else None

    def email_context(self, skip: tuple[str, ...] = ()) -> dict[str, str]:
        """The row as the email writer sees it: your columns, without the agent's bookkeeping."""
        cols = {self.sheet.columns[k] for k in (*NOT_FOR_EMAIL, *skip) if k in self.sheet.columns}
        return self.sheet.row_context(self.index, skip_cols=cols)

    def all_text(self) -> str:
        return " ".join(self.context().values())

    @property
    def ref(self) -> str:
        return f"{self.sheet.store.name}/{self.sheet.title}/row {self.index}"

    @property
    def kind(self) -> str:
        """'buyer' (gets intro emails, 10-day follow-ups) or 'freight' (gets RFQs, 5-day follow-ups)."""
        record_type = self.text("record_type").lower()
        if record_type.startswith("lead"):
            return "buyer"
        if record_type.startswith(("quote", "rfq")):
            return "freight"
        if category := self.text("category"):
            return category_kind(category)
        return self.sheet.kind


class SheetView:
    def __init__(self, store: "WorkbookStore", ws, kind: str):
        self.store = store
        self.ws = ws
        self.kind = kind
        self.header_row = _detect_header_row(ws)
        self.columns: dict[str, int] = {}
        self.headers: dict[int, str] = {}
        if self.header_row:
            self._map_columns()
        self._initial_last_col = self.last_header_col
        self._initial_last_row = self.last_data_row()

    @property
    def title(self) -> str:
        return self.ws.title

    def _map_columns(self) -> None:
        weak: list[tuple[str, int]] = []
        for c in range(1, self.ws.max_column + 1):
            value = _value(self.ws, self.header_row, c)
            if value is None or not str(value).strip():
                continue
            text = str(value).strip()
            self.headers[c] = text
            key = match_header(text)
            if key and key not in self.columns:
                self.columns[key] = c
            elif not key and (weak_key := match_weak_header(text)):
                weak.append((weak_key, c))
        for key, c in weak:
            self.columns.setdefault(key, c)

    @property
    def last_header_col(self) -> int:
        return max(self.headers) if self.headers else 0

    def has(self, key: str) -> bool:
        return key in self.columns

    def ensure(self, key: str) -> int:
        if key in self.columns:
            return self.columns[key]
        field = FIELD_BY_KEY[key]
        if self.header_row is None:
            self.header_row = 1
        col = self.last_header_col + 1
        cell = self.ws.cell(self.header_row, col, field.header)
        template = self.ws.cell(self.header_row, col - 1) if col > 1 else None
        if template is not None and template.has_style:
            cell._style = copy(template._style)
        else:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        letter = get_column_letter(col)
        if not self.ws.column_dimensions[letter].width:
            self.ws.column_dimensions[letter].width = field.width
        self.headers[col] = field.header
        self.columns[key] = col
        self.store.dirty = True
        log.debug("Added column '%s' to %s/%s", field.header, self.store.name, self.title)
        return col

    def last_data_row(self) -> int:
        if not self.header_row:
            return 0
        cols = list(self.headers) or [1]
        for r in range(self.ws.max_row, self.header_row, -1):
            for c in cols:
                if _value(self.ws, r, c) not in (None, ""):
                    return r
        return self.header_row

    def rows(self) -> Iterator[Row]:
        if not self.header_row:
            return
        for r in range(self.header_row + 1, self.last_data_row() + 1):
            if any(_value(self.ws, r, c) not in (None, "") for c in self.headers):
                yield Row(self, r)

    def get(self, r: int, key: str) -> Any:
        c = self.columns.get(key)
        return None if c is None else _value(self.ws, r, c)

    def set(self, r: int, key: str, value: Any) -> None:
        c = self.ensure(key)
        cell = self.ws.cell(r, c)
        if isinstance(value, datetime):
            value = value.date()
        cell.value = value
        if isinstance(value, date):
            cell.number_format = self.store.date_format
        if key in ("source_url", "website") and isinstance(value, str) and value.startswith("http"):
            cell.hyperlink = value
        self.store.dirty = True

    def append(self, values: dict[str, Any]) -> Row:
        for key in values:
            self.ensure(key)
        r = self.last_data_row() + 1
        prev = r - 1 if r - 1 > (self.header_row or 0) else None
        if prev:
            for c in self.headers:
                src = self.ws._cells.get((prev, c))
                if src is not None and src.has_style:
                    dst = self.ws.cell(r, c)
                    dst.font = copy(src.font)
                    dst.border = copy(src.border)
                    dst.alignment = copy(src.alignment)
                    dst.number_format = src.number_format
        for key, value in values.items():
            if value not in (None, ""):
                self.set(r, key, value)
        return Row(self, r)

    def row_context(self, r: int, skip_cols: set[int] | None = None) -> dict[str, str]:
        out: dict[str, str] = {}
        for c, header in self.headers.items():
            if skip_cols and c in skip_cols:
                continue
            value = _value(self.ws, r, c)
            if isinstance(value, (date, datetime)):
                text = fmt_date(value if isinstance(value, date) else value.date())
            else:
                text = clean(value)
            if text:
                out[header] = text
        return out

    def finalize(self) -> None:
        """Stretch Excel tables / filters over any rows and columns the agent added."""
        if not self.header_row:
            return
        last_col = self.last_header_col
        last_row = max(self.last_data_row(), self.header_row + 1)
        for table in list(self.ws.tables.values()):
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if min_row != self.header_row:
                continue
            new_max_col = last_col if max_col >= self._initial_last_col else max_col
            new_max_row = max(max_row, last_row) if max_row >= self._initial_last_row else max_row
            if (new_max_col, new_max_row) == (max_col, max_row):
                continue
            next_id = max((tc.id for tc in table.tableColumns), default=0) + 1
            for c in range(max_col + 1, new_max_col + 1):
                table.tableColumns.append(TableColumn(id=next_id, name=str(self.ws.cell(self.header_row, c).value)))
                next_id += 1
            ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(new_max_col)}{new_max_row}"
            table.ref = ref
            if table.autoFilter is not None:
                table.autoFilter.ref = ref
        if self.ws.auto_filter.ref:
            min_col, min_row, max_col, max_row = range_boundaries(self.ws.auto_filter.ref)
            if min_row == self.header_row:
                self.ws.auto_filter.ref = (
                    f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max(max_col, last_col))}"
                    f"{max(max_row, last_row)}"
                )


class WorkbookStore:
    def __init__(self, name: str, path: Path, cfg):
        self.name = name
        self.path = Path(path)
        self.cfg = cfg
        self.default_kind = cfg["workbooks"][name].get("default_kind", "freight")
        self.date_format = cfg["excel"]["date_format"]
        self.wb = None
        self.dirty = False
        self.load_warnings: list[str] = []
        self._views: dict[str, SheetView] = {}

    def load(self) -> "WorkbookStore":
        if not self.path.exists():
            raise FileNotFoundError(f"Workbook not found: {self.path}")
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            self.wb = openpyxl.load_workbook(self.path, rich_text=True)
        # e.g. "Data Validation extension is not supported and will be removed"
        self.load_warnings = sorted({str(w.message) for w in caught})
        self.dirty = False
        self._views = {}
        return self

    def kind_of(self, title: str) -> str:
        if title in self.cfg.lead_tabs(self.name):
            return self.default_kind
        if category_kind(title) == "freight":
            return "freight"
        if any(word in title.lower() for word in ("roaster", "buyer", "cafe", "café", "importer", "wholesal", "trader")):
            return "buyer"
        return self.default_kind

    def view(self, title: str) -> SheetView | None:
        if title in self._views:
            return self._views[title]
        if title not in self.wb.sheetnames or title in self.cfg["tabs"]["ignore_tabs"]:
            return None
        ws = self.wb[title]
        if not hasattr(ws, "tables"):  # chartsheet
            return None
        self._views[title] = SheetView(self, ws, self.kind_of(title))
        return self._views[title]

    def views(self) -> list[SheetView]:
        return [v for title in self.wb.sheetnames if (v := self.view(title)) is not None]

    def get_or_create(self, title: str, kind: str) -> SheetView:
        if (existing := self.view(title)) is not None:
            return existing
        ws = self.wb.create_sheet(title)
        for c, key in enumerate(NEW_TAB_KEYS[kind], start=1):
            field = FIELD_BY_KEY[key]
            cell = ws.cell(1, c, field.header)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            ws.column_dimensions[get_column_letter(c)].width = field.width
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(NEW_TAB_KEYS[kind]))}1"
        self.dirty = True
        log.info("Created tab '%s' in %s", title, self.path.name)
        return self.view(title)

    def save(self, backup_dir: Path, keep: int = 40) -> bool:
        if not self.dirty:
            return False
        for view in self._views.values():
            view.finalize()
        if reason := excel_lock_reason(self.path):
            raise WorkbookLocked(f"Can't save {self.path.name}: {reason}")
        backup_dir = backup_dir / self.path.stem
        backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.copy2(self.path, backup_dir / f"{self.path.stem}-{stamp}{self.path.suffix}")
        for old in sorted(backup_dir.glob(f"{self.path.stem}-*{self.path.suffix}"))[:-keep]:
            old.unlink(missing_ok=True)
        # "~$" files are skipped by OneDrive sync, so the temp file never uploads.
        tmp = self.path.with_name("~$agent_" + self.path.name)
        self.wb.save(tmp)
        for attempt in range(5):
            try:
                os.replace(tmp, self.path)
                break
            except PermissionError:
                if attempt == 4:
                    tmp.unlink(missing_ok=True)
                    raise WorkbookLocked(f"Can't replace {self.path.name}: file is in use")
                time.sleep(3)
        self.dirty = False
        log.info("Saved %s (backup in %s)", self.path.name, backup_dir)
        return True


def open_stores(cfg, names: list[str]) -> dict[str, WorkbookStore]:
    return {name: WorkbookStore(name, cfg.workbook_path(name), cfg).load() for name in names}


def build_dedupe_index(stores: dict[str, WorkbookStore]) -> DedupeIndex:
    index = DedupeIndex()
    for store in stores.values():
        for view in store.views():
            for row in view.rows():
                company = row.text("company")
                label = f"{store.path.name} > {view.title} row {row.index}" + (f" ({company})" if company else "")
                index.add(label, company=company, website=row.text("website"), email=row.text("email"),
                          phone=row.text("phone"), extra_text=row.all_text())
    return index
